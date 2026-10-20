"""
고시·K-REACH 원본 파일 파서.

모든 고시 별표를 같은 형태의 '스냅샷'으로 바꾼다:
    {(고유번호, 소번호): Entry}
Entry = {uid, sub, name, cas:[...], hazard:{급성/만성/생태: %}, pcts:[...], main, min, deleted}

- PDF 별표(2023-21 ~ 2025-2): pdfplumber 표 추출. 쪽을 넘어가는 행(고유번호 빈칸)은
  직전 항목에 이어 붙인다. 함량은 명칭 문장 안의 "N%" 값에서 읽는다.
- xlsx 별표(2025-19 ~): 병합셀을 풀어서 읽는다. 소번호 병합 구간의 예시 행은
  해당 소번호에 귀속된다(예: 납 화합물 예시 → 97-1-9 소번호 3).
- 개정문 PDF(2022-63, 2022-80): 본문에서 개정·신설된 고유번호 목록만 읽는다.
"""
import html
import re
from collections import OrderedDict

import openpyxl
import pdfplumber

CAS_RE = re.compile(r'\b(\d{2,7}-\d{2}-\d)\b')
UID_RE = re.compile(r'^\d{2,4}-\d-\d{1,4}$')
PCT_RE = re.compile(r'(\d+(?:\.\d+)?)\s*%')
EMBEDDED_UID_RE = re.compile(r'(?:(?<=\s)|^)(\d{2,4}-\d-\d{1,4})(?=\s)')
CATS = ('급성', '만성', '생태')


def serial(uid):
    """고유번호의 일련번호(마지막 숫자). 연도 부분은 재고시 때 바뀔 수 있으나(2023-1-1174 →
    2024-1-1174) 일련번호는 전체에서 유일하므로 고시 간 비교 키로 쓴다."""
    return int(str(uid).rsplit('-', 1)[1])


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    m = re.match(r'\s*(\d+(?:\.\d+)?)', str(v or ''))
    return float(m.group(1)) if m else None


def _cas_list(v):
    return list(OrderedDict.fromkeys(CAS_RE.findall(str(v or ''))))


def _entry(uid, sub, name, cas, hazard, pcts, deleted=False):
    vals = [p for p in pcts if p is not None]
    return {
        'uid': uid, 'sub': sub, 'name': (name or '').strip(), 'cas': cas,
        'hazard': hazard, 'pcts': vals,
        'main': vals[0] if vals else None,
        'min': min(vals) if vals else None,
        'deleted': deleted,
    }


# ── PDF 전체목록 ──────────────────────────────────────────
def parse_list_pdf(path):
    rows = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                for r in table:
                    if not r or len(r) < 2:
                        continue
                    # 좁은 칸에서 고유번호가 줄바꿈되어 추출되기도 한다(예: "2017-1-76\n4")
                    uid = re.sub(r'\s+', '', r[0] or '')
                    # 일부 쪽은 빈 열이 끼어 추출된다 → 비어 있지 않은 나머지 칸을 명칭으로 본다
                    text = '\n'.join(c.strip() for c in r[1:] if c and c.strip())
                    if uid == '고유번호':
                        continue
                    if uid and UID_RE.match(uid):
                        rows.append([uid, text])
                    elif rows:
                        rows[-1][1] += '\n' + text
    # 괘선이 빠진 구간은 두 행이 한 칸으로 합쳐져 추출된다 → 명칭 안의 고유번호에서 다시 나눈다
    split_rows = []
    for uid, text in rows:
        parts = EMBEDDED_UID_RE.split(text)
        split_rows.append([uid, parts[0]])
        for i in range(1, len(parts), 2):
            split_rows.append([parts[i], parts[i + 1]])
    snap = OrderedDict()
    for uid, text in split_rows:
        flat = re.sub(r'\s*\n\s*', ' ', text).strip()
        deleted = flat == '삭제'
        pcts = [float(x) for x in PCT_RE.findall(flat)]
        snap[(uid, None)] = _entry(uid, None, flat, _cas_list(flat), {}, pcts, deleted)
    return snap


# ── 개정문 PDF: 개정·신설 고유번호 ─────────────────────────
def parse_amendment_pdf(path):
    with pdfplumber.open(path) as pdf:
        text = '\n'.join(p.extract_text() or '' for p in pdf.pages)
    head = re.search(r'【별표】\s*중(.*?)다음과\s*같이\s*각각\s*신설', text, re.S)
    if not head:
        raise ValueError(f'개정문 본문을 찾지 못함: {path}')
    body = re.sub(r'\s+', '', head.group(1))
    quoted = re.findall(r'“(\d{2,4}-\d-\d{1,4})”란', body)
    rng = re.search(r'“(\d{2,4}-\d-\d{1,4})”란부터“(\d{2,4}-\d-\d{1,4})”란까지', body)
    amended, created = [], []
    if rng:
        a, b = rng.group(1), rng.group(2)
        pa, pb = a.rsplit('-', 1), b.rsplit('-', 1)
        created = [f'{pa[0]}-{n}' for n in range(int(pa[1]), int(pb[1]) + 1)]
        # 범위 앞에 오는 "“X”란 다음에"의 X는 기준점일 뿐 개정 대상이 아니다
        anchor = re.search(r'“(\d{2,4}-\d-\d{1,4})”란다음에“' + re.escape(a), body)
        skip = {a, b} | ({anchor.group(1)} if anchor else set())
        amended = [u for u in quoted if u not in skip]
    else:
        amended = quoted
    return {'amended': list(OrderedDict.fromkeys(amended)), 'created': created}


# ── xlsx 별표(전체목록·개정사항) ───────────────────────────
def _grid(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:
        v = grid[rng.min_row - 1][rng.min_col - 1]
        for rr in range(rng.min_row - 1, rng.max_row):
            for cc in range(rng.min_col - 1, rng.max_col):
                grid[rr][cc] = v
    return grid


def parse_list_xlsx(path):
    """반환: 스냅샷. 같은 (고유번호, 소번호) 아래 여러 CAS 행은 하나로 합친다.
    종전 유독물질 함량(10번째 열)이 있으면 entry['prev']에 담는다."""
    grid = _grid(path)
    snap = OrderedDict()
    cur = None
    for r in grid[4:]:
        r = (r + [None] * 10)[:10]
        if not any(v not in (None, '') for v in r):
            continue
        uid, sub, ko, en, cas, a, c, e, kind, prev = r
        if uid not in (None, '') and UID_RE.match(str(uid).strip()):
            cur = str(uid).strip()
        if cur is None:
            continue
        sub = int(sub) if _num(sub) is not None else None
        hazard = {k: _num(v) for k, v in zip(CATS, (a, c, e)) if _num(v) is not None}
        key = (cur, sub)
        deleted = str(kind or '').strip() == '삭제'
        if key not in snap:
            snap[key] = _entry(cur, sub, ko, [], hazard, list(hazard.values()), deleted)
            snap[key]['name_en'] = (en or '').strip()
            snap[key]['prev'] = _num(prev)
        ent = snap[key]
        for x in _cas_list(cas):
            if x not in ent['cas']:
                ent['cas'].append(x)
        if not ent['hazard'] and hazard:
            ent['hazard'] = hazard
            ent['pcts'] = list(hazard.values())
            ent['main'] = ent['pcts'][0]
            ent['min'] = min(ent['pcts'])
        if ent.get('prev') is None and _num(prev) is not None:
            ent['prev'] = _num(prev)
    return snap


# ── K-REACH 다운로드 엑셀 ─────────────────────────────────
KREACH_TOKEN = re.compile(
    r'^((?:인체급성|인체만성|생태)(?:·(?:인체급성|인체만성|생태))*)'
    r'(\d{2,4}-\d-\d{1,4})(\(신규화학물질\))?(?:\((\d+)\))?(\([^)]*\))?$')
KREACH_HZ = re.compile(r'(인체급성유해성|인체만성유해성|생태유해성)\s*:\s*([\d.]+)\s*%')


def parse_kreach_xlsx(path):
    """K-REACH 화학물질정보처리시스템 다운로드 → [{cas, name, ke, links:[{uid, sub, cats}], hazard}]"""
    wb = openpyxl.load_workbook(path)
    rows = list(wb.active.iter_rows(values_only=True))
    head = rows[0]
    out = []
    for r in rows[1:]:
        d = dict(zip(head, r))
        links = []
        for tok in html.unescape(str(d.get('급성·만성·생태') or '')).split(','):
            tok = re.sub(r'\s+\(', '(', tok.strip())
            if not tok:
                continue
            m = KREACH_TOKEN.match(tok)
            if not m:
                raise ValueError(f'K-REACH 고유번호 표기 해석 실패: {tok!r}')
            links.append({
                'uid': m.group(2),
                'sub': int(m.group(4)) if m.group(4) else None,
                'cats': [c.replace('인체', '') for c in m.group(1).split('·')],
            })
        hazard = {}
        for seg in str(d.get('유해특성분류 및 혼합물 함량기준(%)') or '').split(' / '):
            if seg.strip().startswith('유독물질'):
                for k, v in KREACH_HZ.findall(seg):
                    cat = k.replace('인체', '').replace('유해성', '')
                    hazard[cat] = min(float(v), hazard.get(cat, float(v)))
        out.append({
            'cas': str(d.get('CAS번호') or '').strip(),
            'name': (d.get('국문명') or d.get('영문명') or '').strip(),
            'ke': str(d.get('기존') or '').strip(),
            'links': links,
            'hazard': hazard,
        })
    return out
